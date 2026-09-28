# SPDX-FileCopyrightText: 2026 Raincloud Maintainers
# SPDX-License-Identifier: Apache-2.0
"""Parse an .xlsx file into one canonical Arrow table, concatenating all sheets.

UCI's Online Retail II ships as a two-sheet Excel file (Year 2009-2010 /
Year 2010-2011). Both sheets share the same 8-column schema. We read every
sheet with openpyxl, add a `sheet_name` column so rows remain traceable to their
source, and union the result. The first row of each sheet is its header, and
every sheet must carry the same header.

Each column's type is inferred from every cell of every sheet, so no cell is
coerced or nulled to fit a guess:
  - only integral numbers -> int64
  - numbers -> float64 (Excel stores IEEE doubles: the sheet XML carries text
    like `2.5499999999999998`, so a decimal type would invent precision)
  - only datetimes / times / durations -> timestamp[us] / time64[us] / duration[us]
  - only booleans -> bool
  - text, alone or beside numbers (Online Retail's `Invoice` holds `489449`
    and the credit note `'C489449'`) -> string, a number written as its exact
    integer text or its shortest round-trip float text
Any other mix (a datetime beside a number, a boolean beside text) has no type
that holds both faithfully and fails the build, naming the first cell of each
kind. Empty cells are null. A formula cell fails too: its cached result may be
absent or stale, and reading it would present a guess as upstream data.

`params.sheet` (optional) restricts to a single named sheet when set.
"""
from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl
import pyarrow as pa
from openpyxl.utils import get_column_letter

_FLOAT_EXACT_INT = 2 ** 53
_TYPES = {"int": pa.int64(), "float": pa.float64(), "str": pa.string(), "bool": pa.bool_(),
          "datetime": pa.timestamp("us"), "time": pa.time64("us"), "duration": pa.duration("us")}


def _kind(value) -> str:
    if isinstance(value, bool): return "bool"
    if isinstance(value, int): return "int"
    if isinstance(value, float): return "int" if value.is_integer() else "float"
    if isinstance(value, str): return "str"
    if isinstance(value, datetime.datetime): return "datetime"
    if isinstance(value, datetime.time): return "time"
    if isinstance(value, datetime.timedelta): return "duration"
    raise TypeError(f"xlsx_parse: unexpected {type(value).__name__} cell value {value!r}")


def _column(name: str, values: list, first: dict[str, str]) -> pa.Array:
    """Build one column in the narrowest type that holds every cell unchanged."""
    kinds = set(first)
    if kinds <= {"int", "float"} and "float" in kinds:
        for value in values:
            if isinstance(value, int) and abs(value) > _FLOAT_EXACT_INT:
                raise ValueError(f"xlsx_parse: column {name!r} mixes fractional numbers with "
                                 f"the integer {value}, which float64 cannot hold exactly")
        return pa.array(values, type=pa.float64())
    if kinds <= {"int", "float", "str"} and "str" in kinds:
        return pa.array([v if v is None or isinstance(v, str)
                         else str(int(v)) if _kind(v) == "int" else repr(v)
                         for v in values], type=pa.string())
    if len(kinds) > 1:
        cells = ", ".join(f"{kind} at {cell}" for kind, cell in sorted(first.items()))
        raise ValueError(f"xlsx_parse: column {name!r} mixes cell kinds with no single "
                         f"faithful type ({cells})")
    kind = kinds.pop() if kinds else "str"
    if kind == "int":
        values = [v if v is None else int(v) for v in values]
    return pa.array(values, type=_TYPES[kind])


def xlsx_parse(spec: dict, parsed: list[tuple[Path, pa.Table | None]], *,
               sheet: str | None = None, **kwargs
               ) -> list[tuple[str, pa.Table]]:
    xlsx_paths = [p for p, _ in parsed if p.suffix.lower() == ".xlsx"]
    if not xlsx_paths:
        raise ValueError("xlsx_parse: no .xlsx file in parsed input")
    path = xlsx_paths[0]

    # data_only=False so a formula shows up as one (data_type "f") instead of
    # as a cached value that may be missing.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            raise ValueError(f"xlsx_parse: {path.name} has no sheet {sheet!r} "
                             f"(sheets: {workbook.sheetnames})")
        header: list[str] | None = None
        sheet_names: list[str] = []
        columns: list[list] = []
        first: list[dict[str, str]] = []
        for name in [sheet] if sheet is not None else workbook.sheetnames:
            where = f"xlsx_parse: {path.name}[{name!r}]"
            rows = workbook[name].iter_rows()
            head = [cell.value for cell in next(rows, ())]
            while head and head[-1] is None: head.pop()
            if header is None:
                if (not head or not all(isinstance(h, str) and h for h in head)
                        or len(set(head)) != len(head) or "sheet_name" in head):
                    raise ValueError(f"{where} header row must be unique non-empty text "
                                     f"other than 'sheet_name', got {head!r}")
                header = head
                columns = [[] for _ in header]
                first = [{} for _ in header]
            elif head != header:
                raise ValueError(f"{where} header {head!r} differs from the first "
                                 f"sheet's {header!r}")
            count, blank = 0, None
            for number, row in enumerate(rows, 2):
                values = [cell.value for cell in row]
                if all(v is None for v in values):
                    blank = blank or number
                    continue
                # Trailing empty rows are sheet formatting; an empty row with
                # data after it is a gap whose meaning only a recipe can state.
                if blank is not None:
                    raise ValueError(f"{where} row {blank} is empty between data rows")
                if any(v is not None for v in values[len(header):]):
                    raise ValueError(f"{where} row {number} has cells beyond the "
                                     f"{len(header)}-column header")
                values += [None] * (len(header) - len(values))
                for i, (cell, value) in enumerate(zip(row, values)):
                    if value is None: continue
                    ref = f"{get_column_letter(i + 1)}{number}"
                    if cell.data_type == "f":
                        raise ValueError(f"{where} cell {ref} is a formula ({value!r}); "
                                         f"its cached result is not upstream data")
                    first[i].setdefault(_kind(value), f"{name!r}!{ref}")
                for column, value in zip(columns, values):
                    column.append(value)
                count += 1
            sheet_names.extend([name] * count)
            print(f"  {path.name}[{name!r}]: {count:,} rows × {len(header) + 1} cols")
    finally:
        workbook.close()

    merged = pa.table({"sheet_name": pa.array(sheet_names, type=pa.string()),
                       **{h: _column(h, c, f) for h, c, f in zip(header, columns, first)}})
    print(f"  merged: {merged.num_rows:,} rows × {merged.num_columns} cols")
    return [(spec["slug"], merged)]
